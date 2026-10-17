"""Compare API verdicts and conditional localization accuracy against open_labels.xlsx."""
import math
from pathlib import Path
import re
import sys

try:
    from .common import ToolError, analyze, parser, positive
except ImportError:
    from common import ToolError, analyze, parser, positive


def parse_ts(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value < 0:
            raise ValueError("invalid label timestamp")
        return float(value)
    text = str(value).strip()
    minute = re.search(r"(\d+)\s*минут", text)
    second = re.search(r"(\d+)\s*секунд", text)
    if not minute and not second:
        raise ValueError("unrecognized label timestamp")
    return float((int(minute.group(1)) * 60 if minute else 0) + (int(second.group(1)) if second else 0))


def video_number(name):
    match = re.search(r"(\d+)", str(name))
    if not match:
        raise ValueError("video name has no number")
    return int(match.group(1))


def load_labels(xlsx):
    from openpyxl import load_workbook
    workbook = load_workbook(xlsx, read_only=True)
    labels = {}
    try:
        for row in workbook.active.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            number = video_number(row[0])
            if number in labels:
                raise ValueError("duplicate video number in labels")
            label = str(row[2]).strip().lower()
            if label not in {"нарушение", "нет нарушения", "без нарушения", "не нарушение"}:
                raise ValueError("unrecognized verdict label")
            labels[number] = ("violation" if label == "нарушение" else "no_violation", parse_ts(row[1]))
    finally:
        workbook.close()
    if not labels:
        raise ValueError("no labels found")
    return labels


def main(argv=None):
    cli = parser(__doc__)
    cli.add_argument("root", type=Path, help="folder containing videos/ and open_labels.xlsx")
    cli.add_argument("tolerance_sec", nargs="?", type=positive, default=3.0)
    args = cli.parse_args(argv)
    try:
        labels = load_labels(args.root / "open_labels.xlsx")
        files = {}
        for path in (args.root / "videos").iterdir():
            if path.suffix.lower() != ".mov":
                continue
            number = video_number(path.stem)
            if number in files:
                raise ValueError("duplicate video number in videos folder")
            files[number] = path
        missing = set(labels) - set(files)
        if missing:
            raise ToolError(f"missing {len(missing)} labelled videos; evaluation would be incomplete")
        correct = localized = localization_total = 0
        slowest = 0.0
        for number, (truth, timestamp) in sorted(labels.items()):
            job, duration = analyze(args, files[number])
            slowest = max(slowest, duration)
            prediction = job["verdict"]
            correct += prediction == truth
            near = None
            if truth == "violation" and prediction == "violation":
                if timestamp is None:
                    raise ValueError("violation label has no timestamp")
                localization_total += 1
                near = any(abs(v["timecode_sec"] - timestamp) <= args.tolerance_sec for v in job["violations"])
                localized += near
            print(f"video {number}: true={truth} pred={prediction} time_ok={near} {duration:.3f}s")
        print(f"verdict accuracy: {correct}/{len(labels)}; timecode within {args.tolerance_sec}s: {localized}/{localization_total}; slowest: {slowest:.3f}s")
        return 1 if slowest >= 300 else 0
    except ToolError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    except (OSError, ValueError, KeyError, TypeError):
        print("ERROR: invalid/missing dataset, label workbook, or API result; check inputs", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
